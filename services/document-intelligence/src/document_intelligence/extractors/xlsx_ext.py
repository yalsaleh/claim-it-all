from io import BytesIO

from openpyxl import load_workbook

from document_intelligence.config import get_settings
from document_intelligence.extractors.base import ExtractionResult, SegmentDraft


def extract_xlsx(data: bytes) -> ExtractionResult:
    settings = get_settings()
    try:
        workbook = load_workbook(BytesIO(data), data_only=False, read_only=True)
    except Exception as exc:  # noqa: BLE001
        return ExtractionResult(
            status="FAILED",
            warnings=[f"xlsx_parse_error:{exc.__class__.__name__}"],
            content_may_be_incomplete=True,
        )

    segments: list[SegmentDraft] = []
    warnings: list[str] = []
    cell_count = 0
    texts: list[str] = []

    for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
        rows_out: list[str] = []
        for row in sheet.iter_rows(values_only=False):
            values = []
            for cell in row:
                cell_count += 1
                if cell_count > settings.max_workbook_cells:
                    warnings.append("workbook_cell_limit_reached")
                    return ExtractionResult(
                        status="PARTIALLY_SUCCEEDED",
                        warnings=warnings,
                        metadata={"cellsSeen": cell_count, "sheetCount": len(workbook.sheetnames)},
                        plain_text="\n".join(texts),
                        segments=segments,
                        content_may_be_incomplete=True,
                    )
                if cell.value is None:
                    continue
                if getattr(cell, "data_type", None) == "f":
                    values.append(f"={cell.value}")
                else:
                    values.append(str(cell.value))
            if values:
                rows_out.append("\t".join(values))
        sheet_text = "\n".join(rows_out)
        texts.append(f"# {sheet.title}\n{sheet_text}")
        segments.append(
            SegmentDraft(
                kind="SPREADSHEET_SHEET",
                ordinal=sheet_index,
                label=sheet.title,
                text=sheet_text[:20000] if sheet_text else None,
                locator={
                    "schemaVersion": 1,
                    "kind": "spreadsheet_sheet",
                    "sheetName": sheet.title,
                },
            )
        )

    return ExtractionResult(
        status="SUCCEEDED",
        warnings=warnings,
        metadata={"sheetCount": len(workbook.sheetnames), "cellsSeen": cell_count},
        plain_text="\n\n".join(texts),
        segments=segments,
        ocr_used=False,
        content_may_be_incomplete=False,
    )
