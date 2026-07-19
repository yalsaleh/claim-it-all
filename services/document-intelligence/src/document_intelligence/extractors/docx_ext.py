from io import BytesIO

from docx import Document

from document_intelligence.extractors.base import ExtractionResult, SegmentDraft


def extract_docx(data: bytes) -> ExtractionResult:
    try:
        document = Document(BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        return ExtractionResult(
            status="FAILED",
            warnings=[f"docx_parse_error:{exc.__class__.__name__}"],
            content_may_be_incomplete=True,
        )

    segments: list[SegmentDraft] = []
    texts: list[str] = []
    ordinal = 0
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        ordinal += 1
        texts.append(text)
        segments.append(
            SegmentDraft(
                kind="PARAGRAPH",
                ordinal=ordinal,
                label=None,
                text=text,
                locator={"schemaVersion": 1, "kind": "paragraph", "label": f"p{ordinal}"},
            )
        )

    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append(" | ".join(cell.text.strip() for cell in row.cells))
        table_text = "\n".join(rows)
        if table_text.strip():
            ordinal += 1
            texts.append(table_text)
            segments.append(
                SegmentDraft(
                    kind="SECTION",
                    ordinal=ordinal,
                    label=f"Table {table_index}",
                    text=table_text,
                    locator={
                        "schemaVersion": 1,
                        "kind": "section",
                        "label": f"table-{table_index}",
                    },
                )
            )

    return ExtractionResult(
        status="SUCCEEDED" if texts else "PARTIALLY_SUCCEEDED",
        warnings=[] if texts else ["docx_empty"],
        metadata={"paragraphAndTableSegments": len(segments)},
        plain_text="\n\n".join(texts),
        segments=segments,
        ocr_used=False,
        content_may_be_incomplete=False,
    )
